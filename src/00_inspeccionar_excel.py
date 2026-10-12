"""Muestra la estructura de los Excel: hojas, número de filas y las primeras filas de cada hoja.

Uso:  python src/00_inspeccionar_excel.py                      (todos los .xlsx bajo data/)
      python src/00_inspeccionar_excel.py ruta/archivo.xlsx    (uno en particular)
Copia la salida completa y compártela para escribir el lector de la fuente B (leer_b en job1).
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

import config as C

FILAS_MUESTRA = 8
COLS_MUESTRA = 15

archivos = [Path(a) for a in sys.argv[1:]] or sorted(C.DATA.rglob("*.xlsx"))
if not archivos:
    sys.exit("No hay archivos .xlsx en data/. Los .xls antiguos deben abrirse y guardarse como .xlsx.")

for archivo in archivos:
    print("=" * 90)
    print(f"ARCHIVO: {archivo}  ({archivo.stat().st_size / 1_048_576:.1f} MiB)")
    wb = load_workbook(archivo, read_only=True, data_only=True)
    for ws in wb.worksheets:
        filas, muestra, max_cols = 0, [], 0
        for fila in ws.iter_rows(values_only=True):
            filas += 1
            max_cols = max(max_cols, len(fila))
            if filas <= FILAS_MUESTRA:
                muestra.append(fila)
        print(f"\n--- Hoja: '{ws.title}' | filas: {filas} | columnas: {max_cols}")
        for fila in muestra:
            extra = f" ... (+{len(fila) - COLS_MUESTRA} columnas)" if len(fila) > COLS_MUESTRA else ""
            print(fila[:COLS_MUESTRA], extra)
    wb.close()
